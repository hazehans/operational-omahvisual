import openpyxl
import pandas as pd

wb1 = openpyxl.load_workbook('raw_data/DATA TOOLS BOX .2 (1).xlsx', data_only=True)
wb2 = openpyxl.load_workbook('raw_data/Salinan dari DEV_STOCK OPNAME by Risma.xlsx', data_only=True)

with open('deep_analysis.txt', 'w', encoding='utf-8') as out:
    out.write("=== FILE 1: ALMARI 1, 2, 3 DETAILS ===\n")
    for sname in ['STOCK ALMARI 1', 'STOCK ALMARI 2', 'ALMARI 3']:
        ws = wb1[sname]
        out.write(f"\n--- {sname} ---\n")
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 10))]
            if any(v is not None for v in row_vals):
                out.write(f"R{r}: {[v for v in row_vals if v is not None]}\n")

    out.write("\n\n=== FILE 2: EVOLUTION OF STOCK OPNAME SHEETS ===\n")
    # check row counts of stock sheets
    sheets_to_compare = [
        'INVENTORY ASET LED 13 NOV 25',
        'INVENTORY ASET LED 18 DES 2025',
        'FIX STOCK',
        ' STOCK',
        'STOCK - AURA&PUTRI',
        'MASTER DATA 10 APR 2026',
        'Burhan',
        'MASTER DATA mei 2026'
    ]
    for s in sheets_to_compare:
        ws = wb2[s]
        # find title row and header
        title = ws.cell(3, 2).value or ws.cell(1, 2).value or ws.cell(7, 4).value
        out.write(f"Sheet '{s}': Title='{title}', MaxRow={ws.max_row}, MaxCol={ws.max_column}\n")

    out.write("\n\n=== FILE 2: HARDCASE TV & ASSETS ===\n")
    ws_hc = wb2['UKURAN HARDCASE TV']
    for r in range(1, ws_hc.max_row + 1):
        vals = [ws_hc.cell(r, c).value for c in range(1, ws_hc.max_column + 1)]
        if any(v is not None for v in vals):
            out.write(f"HC R{r}: {[v for v in vals if v is not None]}\n")

    out.write("\n\n=== FILE 2: TAMBAHAN ASET (Nov 2025) ===\n")
    ws_ta = wb2['TAMBAHAN ASET']
    for r in range(1, ws_ta.max_row + 1):
        vals = [ws_ta.cell(r, c).value for c in range(1, ws_ta.max_column + 1)]
        if any(v is not None for v in vals):
            out.write(f"TA R{r}: {[v for v in vals if v is not None]}\n")

print("Deep analysis completed!")

