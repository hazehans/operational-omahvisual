import openpyxl
import os

wb_path = 'raw_data/Salinan dari DEV_STOCK OPNAME by Risma.xlsx'
wb = openpyxl.load_workbook(wb_path, data_only=True)

with open('file2_detailed_sheets.txt', 'w', encoding='utf-8') as out:
    for name in wb.sheetnames:
        ws = wb[name]
        out.write('='*60 + '\n')
        out.write(f'SHEET: "{name}" (Max Row: {ws.max_row}, Max Col: {ws.max_column})\n')
        
        found = 0
        for r in range(1, min(ws.max_row + 1, 20)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 25))]
            if any(v is not None for v in row_vals):
                non_empty = [f'C{c}:{v}' for c, v in enumerate(row_vals, 1) if v is not None]
                out.write(f'  Row {r}: ' + ' | '.join(non_empty[:10]) + '\n')
                found += 1
                if found >= 6:
                    break

print("Done writing file2_detailed_sheets.txt")

