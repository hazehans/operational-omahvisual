import openpyxl
import pandas as pd

file2_path = 'raw_data/Salinan dari DEV_STOCK OPNAME by Risma.xlsx'
wb2 = openpyxl.load_workbook(file2_path, data_only=True)

# 1. Analyze "Data Bersih"
ws_db = wb2['Data Bersih']
data_bersih_rows = []
for r in range(2, ws_db.max_row + 1):
    vals = [ws_db.cell(r, c).value for c in range(1, 10)]
    if any(v is not None for v in vals):
        data_bersih_rows.append(vals)

df_db = pd.DataFrame(data_bersih_rows, columns=[
    'kategori_utama', 'sub_kategori', 'merk_seri', 'nama_barang', 
    'spesifikasi', 'qty', 'satuan', 'kondisi', 'keterangan'
])

# 2. Analyze TV Master
ws_tv = wb2['INVENTORY ASET TV MASTER']
tv_rows = []
for r in range(7, min(ws_tv.max_row + 1, 100)):
    row_vals = [ws_tv.cell(r, c).value for c in range(2, 9)]
    if any(v is not None for v in row_vals):
        tv_rows.append(row_vals)

# 3. Analyze Kekurangan
ws_kekurangan = wb2['KEKURANGAN']
kekurangan_rows = []
for r in range(4, min(ws_kekurangan.max_row + 1, 25)):
    vals = [ws_kekurangan.cell(r, c).value for c in range(2, 9)]
    if any(v is not None for v in vals):
        kekurangan_rows.append(vals)

with open('key_sheets_analysis.txt', 'w', encoding='utf-8') as out:
    out.write("Data Bersih summary:\n")
    out.write(f"Total rows: {len(df_db)}\n\n")
    out.write("Kategori Utama distribution:\n")
    out.write(df_db['kategori_utama'].value_counts(dropna=False).to_string() + "\n\n")
    out.write("Sub-kategori distribution:\n")
    out.write(df_db['sub_kategori'].value_counts(dropna=False).to_string() + "\n\n")
    out.write("Satuan distribution:\n")
    out.write(df_db['satuan'].value_counts(dropna=False).to_string() + "\n\n")

    out.write("Sample rows of Data Bersih (first 25 rows):\n")
    for idx, r in df_db.head(25).iterrows():
        out.write(f"  [{r['kategori_utama']} > {r['sub_kategori']}] {r['nama_barang']} ({r['merk_seri'] or ''} {r['spesifikasi'] or ''}) - Qty: {r['qty']} {r['satuan']} | Kondisi: {r['kondisi'] or '-'} | Ket: {r['keterangan'] or '-'}\n")

    out.write("\nTV Master sample (first 15 rows):\n")
    for row in tv_rows[:15]:
        out.write(f"  {row}\n")

    out.write("\nKekurangan sample (first 15 rows):\n")
    for row in kekurangan_rows[:15]:
        out.write(f"  {row}\n")

print("Analysis successfully written to key_sheets_analysis.txt")

